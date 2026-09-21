"""
Gerador do Relatorio.docx atualizado a partir do Relatorio_Executivo.md
"""
import datetime
from docx import Document
from docx.shared import Pt

md_path  = r"reports/Relatorio_Executivo.md"
out_path = r"reports/Relatorio.docx"

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()

# Fonte padrão
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

for raw in lines:
    line = raw.rstrip()
    if line.startswith("# "):
        doc.add_heading(line[2:], level=1)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=2)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=3)
    elif line.startswith("* ") or line.startswith("- "):
        doc.add_paragraph(line[2:], style="List Bullet")
    elif line.strip() in ("", "---"):
        doc.add_paragraph("")
    else:
        doc.add_paragraph(line)

# Rodapé
section = doc.sections[0]
footer_text = "PC Descomplicado — Air Company AI Fellowship — Atualizado em " + \
              datetime.date.today().strftime("%d/%m/%Y")
section.footer.paragraphs[0].text = footer_text

doc.save(out_path)
print("DOCX salvo com sucesso:", out_path)
